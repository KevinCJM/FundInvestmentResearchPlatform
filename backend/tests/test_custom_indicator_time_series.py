from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
for path in (ROOT, BACKEND):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from backend.product_analysis_numba import kdj_kernel  # noqa: E402
from cal_indicators.typed_dsl import compose_typed_series_bundle  # noqa: E402
from cal_indicators.typed_numba_kernels import (  # noqa: E402
    BASIC_OPCODES,
    binary_scalar,
    series_safe_divide_1d,
)
from custom_indicators.errors import ConflictError, ValidationError  # noqa: E402
from custom_indicators.runtime_context import aligned_return_series_kernel  # noqa: E402
from custom_indicators.series_definitions import (  # noqa: E402
    normalize_time_series_definition,
    parameter_variable_types,
    series_expressions,
)
from custom_indicators.service import CustomIndicatorService  # noqa: E402
from custom_indicators.variable_registry import variable_types  # noqa: E402
from services import custom_indicator_routes  # noqa: E402


def _write_market_data(root: Path, count: int = 80) -> pd.DataFrame:
    dates = pd.bdate_range("2025-01-02", periods=count)
    positions = np.arange(count, dtype=np.float64)
    close = 10.0 + positions * 0.04 + np.sin(positions / 4.0) * 0.25
    frame = pd.DataFrame(
        {
            "ts_code": "510300.SH",
            "date": dates,
            "open": close - 0.05,
            "high": close + 0.20,
            "low": close - 0.20,
            "close": close,
            "pre_close": np.r_[close[0], close[:-1]],
            "change": np.r_[0.0, np.diff(close)],
            "pct_chg": np.r_[0.0, np.diff(close) / close[:-1] * 100.0],
            "vol": 1_000.0 + positions * 10.0,
            "amount": (1_000.0 + positions * 10.0) * close,
        }
    )
    frame.to_parquet(root / "etf_daily_candle_df.parquet", index=False)
    adjusted_nav = close / close[0]
    pd.DataFrame(
        {
            "ts_code": "510300.SH",
            "name": "沪深300ETF",
            "date": dates,
            "adj_nav": adjusted_nav,
        }
    ).to_parquet(root / "etf_daily_df.parquet", index=False)
    pd.DataFrame(
        {
            "ts_code": ["510300.SH"],
            "name": ["沪深300ETF"],
            "list_date": ["20120101"],
        }
    ).to_parquet(root / "etf_info_df.parquet", index=False)
    fund_nav = 1.5 + positions * 0.003 + np.sin(positions / 5.0) * 0.015
    pd.DataFrame(
        {
            "ts_code": ["000001.OF"],
            "code": ["000001"],
            "name": ["测试公募基金"],
        }
    ).to_parquet(root / "fund_info_df.parquet", index=False)
    pd.DataFrame(
        {
            "ts_code": "000001.OF",
            "name": "测试公募基金",
            "date": dates,
            "adj_nav": fund_nav,
        }
    ).to_parquet(root / "fund_nav_df.parquet", index=False)
    return frame.assign(adj_nav=adjusted_nav)


def _service(tmp_path: Path) -> tuple[CustomIndicatorService, pd.DataFrame]:
    frame = _write_market_data(tmp_path)
    service = CustomIndicatorService(tmp_path, tmp_path)
    status = service.warm_numba_plans()
    assert status["time_series_plans"] == 5
    return service, frame


def _instance(indicator_id: str) -> dict[str, object]:
    return {"indicator_id": indicator_id}


def _as_float(values: list[float | None]) -> np.ndarray:
    return np.asarray([np.nan if value is None else value for value in values], dtype=np.float64)


def test_array_safe_divide_returns_nan_per_row_without_changing_scalar_diagnostic() -> None:
    numerator = np.ascontiguousarray(np.array([1.0, 2.0, 3.0]))
    denominator = np.ascontiguousarray(np.array([2.0, 0.0, 1e-13]))
    result = series_safe_divide_1d(numerator, denominator)
    assert result == pytest.approx(np.array([0.5, np.nan, np.nan]), nan_ok=True)
    assert series_safe_divide_1d.nopython_signatures
    with pytest.raises(ValueError, match="DIVIDE_BY_ZERO"):
        binary_scalar(BASIC_OPCODES["divide"], 1.0, 0.0)


def test_aligned_return_series_kernel_preserves_axis_and_missing_values() -> None:
    levels = np.ascontiguousarray(
        np.array([1.0, 1.1, np.nan, 1.2, 0.0, 1.3, 1.43, np.inf])
    )
    simple, logarithmic = aligned_return_series_kernel(levels)
    expected_simple = np.array([np.nan, 0.1, np.nan, np.nan, np.nan, np.nan, 0.1, np.nan])
    expected_log = np.array(
        [np.nan, np.log(1.1), np.nan, np.nan, np.nan, np.nan, np.log(1.1), np.nan]
    )
    assert simple.flags.c_contiguous
    assert logarithmic.flags.c_contiguous
    assert simple == pytest.approx(expected_simple, nan_ok=True)
    assert logarithmic == pytest.approx(expected_log, nan_ok=True)
    assert aligned_return_series_kernel.nopython_signatures


def test_five_builtin_time_series_indicators_use_fixed_formulas_and_match_references(
    tmp_path: Path,
) -> None:
    service, frame = _service(tmp_path)
    definitions = {
        item["id"]: item for item in service.list_indicators()["items"]
        if item.get("result_kind") == "time_series"
    }
    assert all(not item.get("parameter_schema") for item in definitions.values())
    assert definitions["builtin-close-moving-average-series"]["series_outputs"][0]["expression"] == "rolling_mean(market_close, 20)"
    assert definitions["builtin-volume-moving-average-series"]["series_outputs"][0]["expression"] == "rolling_mean(volume, 10)"
    rolling_definition = definitions["builtin-rolling-5d-annualized-sharpe-series"]
    assert rolling_definition["series_outputs"][0]["expression"] == (
        "(rolling_mean(returns, 5) - risk_free_rate_per_observation) / "
        "rolling_std(returns, 5, 1) * sqrt(periods_per_year)"
    )
    assert rolling_definition["rolling_source"] == {
        **rolling_definition["rolling_source"],
        "kind": "rolling_scalar",
        "indicator_id": "builtin-annualized-sharpe-v2",
        "indicator_revision": 1,
        "window_observations": 5,
        "minimum_observations": 5,
        "detached": False,
    }

    response = service.evaluate_series(
        indicator_instances=[
            _instance("builtin-close-moving-average-series"),
            _instance("builtin-bollinger-bands-series"),
            _instance("builtin-volume-moving-average-series"),
            _instance("builtin-kdj-series"),
            _instance("builtin-rolling-5d-annualized-sharpe-series"),
        ],
        target={"kind": "etf", "product_id": "510300.SH"},
        period="ALL",
    )

    assert response["summary"] == {
        "total": 5,
        "ok": 5,
        "warning": 0,
        "unavailable": 0,
        "error": 0,
    }
    assert response["execution"]["python_fallback"] == 0
    assert response["execution"]["python_operator_calls"] == 0
    assert response["execution"]["request_time_compilation"] == 0
    by_id = {item["indicator_id"]: item for item in response["results"]}
    close = frame["close"].to_numpy(dtype=np.float64)
    volume = frame["vol"].to_numpy(dtype=np.float64)

    ma = by_id["builtin-close-moving-average-series"]
    expected_ma = pd.Series(close).rolling(20, min_periods=20).mean().to_numpy()
    assert _as_float(ma["channels"][0]["values"]) == pytest.approx(expected_ma, nan_ok=True)

    boll = by_id["builtin-bollinger-bands-series"]
    channels = {item["id"]: item["values"] for item in boll["channels"]}
    middle = pd.Series(close).rolling(20, min_periods=20).mean().to_numpy()
    deviation = pd.Series(close).rolling(20, min_periods=20).std(ddof=0).to_numpy()
    for name, expected in {
        "upper": middle + 2.0 * deviation,
        "middle": middle,
        "lower": middle - 2.0 * deviation,
    }.items():
        assert _as_float(channels[name]) == pytest.approx(expected, nan_ok=True)

    volume_ma = by_id["builtin-volume-moving-average-series"]
    expected_volume = pd.Series(volume).rolling(10, min_periods=10).mean().to_numpy()
    assert _as_float(volume_ma["channels"][0]["values"]) == pytest.approx(expected_volume, nan_ok=True)

    kdj = by_id["builtin-kdj-series"]
    expected_kdj = kdj_kernel(
        np.ascontiguousarray(frame["high"].to_numpy(dtype=np.float64)),
        np.ascontiguousarray(frame["low"].to_numpy(dtype=np.float64)),
        np.ascontiguousarray(close),
        9,
        3,
        3,
    )
    for row, channel in enumerate(kdj["channels"]):
        assert np.asarray(channel["values"], dtype=np.float64) == pytest.approx(expected_kdj[row])

    rolling_sharpe = by_id["builtin-rolling-5d-annualized-sharpe-series"]
    returns = pd.Series(frame["adj_nav"]).pct_change()
    risk_free = (1.0 + 0.015) ** (1.0 / 252.0) - 1.0
    expected_sharpe = (
        (returns.rolling(5, min_periods=5).mean() - risk_free)
        / returns.rolling(5, min_periods=5).std(ddof=1)
        * np.sqrt(252.0)
    ).to_numpy()
    assert _as_float(rolling_sharpe["channels"][0]["values"]) == pytest.approx(
        expected_sharpe,
        nan_ok=True,
    )
    assert rolling_sharpe["lookback_observations"] == 5
    assert rolling_sharpe["minimum_observations"] == 5


def test_rolling_sharpe_supports_public_fund_adjusted_nav(tmp_path: Path) -> None:
    service, _frame = _service(tmp_path)
    response = service.evaluate_series(
        indicator_instances=[
            _instance("builtin-rolling-5d-annualized-sharpe-series")
        ],
        target={"kind": "fund", "product_id": "000001.OF"},
        period="ALL",
    )
    result = response["results"][0]
    assert result["status"] == "ok"
    values = _as_float(result["channels"][0]["values"])
    nav = pd.read_parquet(tmp_path / "fund_nav_df.parquet")["adj_nav"]
    returns = nav.pct_change()
    risk_free = (1.0 + 0.015) ** (1.0 / 252.0) - 1.0
    expected = (
        (returns.rolling(5, min_periods=5).mean() - risk_free)
        / returns.rolling(5, min_periods=5).std(ddof=1)
        * np.sqrt(252.0)
    ).to_numpy()
    assert values == pytest.approx(expected, nan_ok=True)
    assert response["execution"]["python_fallback"] == 0


def test_rolling_sharpe_zero_volatility_is_missing_not_a_full_plan_failure(
    tmp_path: Path,
) -> None:
    _write_market_data(tmp_path)
    nav = pd.read_parquet(tmp_path / "etf_daily_df.parquet")
    nav["adj_nav"] = 1.0
    nav.to_parquet(tmp_path / "etf_daily_df.parquet", index=False)
    service = CustomIndicatorService(tmp_path, tmp_path)
    service.warm_numba_plans()
    response = service.evaluate_series(
        indicator_instances=[
            _instance("builtin-rolling-5d-annualized-sharpe-series")
        ],
        target={"kind": "etf", "product_id": "510300.SH"},
        period="ALL",
    )
    result = response["results"][0]
    assert result["status"] == "unavailable"
    assert result["warnings"] == []
    assert len(result["channels"]) == 1
    assert set(result["channels"][0]["values"]) == {None}
    assert response["execution"]["python_fallback"] == 0
    assert response["execution"]["request_time_compilation"] == 0


def test_rolling_sharpe_is_causal_and_does_not_use_future_nav(tmp_path: Path) -> None:
    left = tmp_path / "left"
    right = tmp_path / "right"
    left.mkdir()
    right.mkdir()
    _write_market_data(left)
    _write_market_data(right)
    changed = pd.read_parquet(right / "etf_daily_df.parquet")
    changed.loc[40:, "adj_nav"] = changed.loc[40:, "adj_nav"] * 1.35
    changed.to_parquet(right / "etf_daily_df.parquet", index=False)

    outputs: list[np.ndarray] = []
    for data_dir in (left, right):
        service = CustomIndicatorService(data_dir, data_dir)
        service.warm_numba_plans()
        response = service.evaluate_series(
            indicator_instances=[
                _instance("builtin-rolling-5d-annualized-sharpe-series")
            ],
            target={"kind": "etf", "product_id": "510300.SH"},
            period="ALL",
        )
        outputs.append(_as_float(response["results"][0]["channels"][0]["values"]))

    assert outputs[0][:40] == pytest.approx(outputs[1][:40], nan_ok=True)
    assert not np.allclose(outputs[0][40:], outputs[1][40:], equal_nan=True)


def test_time_series_plan_shares_bollinger_subexpressions(tmp_path: Path) -> None:
    service, _frame = _service(tmp_path)
    definition = service.indicators.get("builtin-bollinger-bands-series")
    plan = compose_typed_series_bundle(
        series_expressions(definition),
        variable_types={
            **variable_types("single_product", definition["dsl_version"]),
            **parameter_variable_types(definition),
        },
        dsl_version=definition["dsl_version"],
        operator_registry_version=definition["operator_registry_version"],
    )
    operators = [node.operator_id for node in plan.nodes]
    assert operators.count("rolling_mean") == 1
    assert operators.count("rolling_std") == 1
    assert set(plan.roots) == {"upper", "middle", "lower"}


def test_runtime_algorithm_parameters_are_rejected(tmp_path: Path) -> None:
    service, _frame = _service(tmp_path)
    with pytest.raises(ValidationError) as error:
        service.evaluate_series(
            indicator_instances=[{
                "indicator_id": "builtin-close-moving-average-series",
                "parameters": {"window": 5},
            }],
            target={"kind": "etf", "product_id": "510300.SH"},
            period="ALL",
        )
    assert error.value.code in {
        "SERIES_PARAMETERS_FIXED_IN_DEFINITION",
        "SERIES_RUNTIME_PARAMETERS_NOT_SUPPORTED",
        "UNKNOWN_SERIES_PARAMETER",
    }


def test_legacy_runtime_parameter_definition_is_frozen_to_default_value() -> None:
    fields = {
        "name": "旧版 5 日均线",
        "description": "兼容迁移",
        "result_kind": "time_series",
        "output_contract": "series_bundle",
        "expression": "rolling_mean(market_close, window, window)",
        "series_outputs": [{
            "id": "ma",
            "label": "5 日均线",
            "expression": "rolling_mean(market_close, window, window)",
            "unit": "",
            "display_format": "number",
            "precision": 4,
            "output_measure": "auto",
        }],
        "parameter_schema": [{
            "id": "window",
            "label": "周期",
            "type": "integer",
            "default": 5,
            "minimum": 2,
            "maximum": 500,
            "step": 1,
            "description": "",
        }],
        "axis_anchor": "market_close",
        "history_policy": "lookback",
        "lookback_parameter": "window",
    }
    normalized = normalize_time_series_definition(fields)
    assert normalized["parameter_schema"] == []
    assert normalized["series_outputs"][0]["expression"] == "rolling_mean(market_close, 5)"
    assert normalized["fixed_parameters"][0]["id"] == "window"
    assert normalized["fixed_parameters"][0]["value"] == 5


def test_scalar_evaluation_rejects_time_series_definition(tmp_path: Path) -> None:
    service, _frame = _service(tmp_path)
    with pytest.raises(ValidationError) as error:
        service.evaluate(
            indicator_ids=["builtin-kdj-series"],
            inline_definition=None,
            targets=[{"kind": "etf", "product_id": "510300.SH"}],
            period="ALL",
        )
    assert error.value.code == "INDICATOR_RESULT_KIND_MISMATCH"


def test_validation_exposes_true_math_latex_measure_and_inferred_history(tmp_path: Path) -> None:
    service, _frame = _service(tmp_path)
    definition = service.get_indicator("builtin-bollinger-bands-series")
    validation = service.validate(definition)

    assert validation["valid"] is True
    assert set(validation["dag"]["roots"]) == {"upper", "middle", "lower"}
    assert validation["history_policy"] == "lookback"
    assert validation["lookback_observations"] == 20
    assert validation["minimum_observations"] == 20
    for item in validation["output_inferences"].values():
        latex = item["display_latex"]
        assert latex
        assert "rolling\\_mean" not in latex
        assert "rolling_mean" not in latex
        assert r"\begin{cases}" not in latex
        assert r"\mathrm{NaN}" not in latex
        assert item["resolved_output_measure"] == "raw_market_price"
        assert item["semantic_dimension"] == "raw_market_price"
    assert r"\mu_{t,20}" in validation["output_inferences"]["middle"]["display_latex"]
    assert r"\sigma_{t,20}" in validation["output_inferences"]["upper"]["display_latex"]
    assert all("latex_fragment" in node for node in validation["dag"]["nodes"])
    assert "chart_panel" not in definition
    assert "chart_panel" not in definition["presentation"]


def test_rolling_sharpe_validate_route_returns_math_latex(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """A catalog definition must survive the public GET → validate round trip."""

    service, _frame = _service(tmp_path)
    monkeypatch.setattr(custom_indicator_routes, "indicator_service", service)
    app = FastAPI()
    app.include_router(custom_indicator_routes.router)
    definition = service.get_indicator(
        "builtin-rolling-5d-annualized-sharpe-series"
    )

    response = TestClient(app).post(
        "/api/custom-indicators/validate",
        json=definition,
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["valid"] is True
    assert payload["diagnostics"] == []
    output = payload["output_inferences"]["value"]
    assert output["display_latex"]
    assert "rolling_mean" not in output["display_latex"]
    assert "rolling_std" not in output["display_latex"]
    assert payload["dag"]["roots"]["value"] == output["root_id"]


def test_kdj_measure_inference_distinguishes_bounded_kd_from_unbounded_j(tmp_path: Path) -> None:
    service, _frame = _service(tmp_path)
    validation = service.validate(service.get_indicator("builtin-kdj-series"))
    assert validation["valid"] is True
    output = validation["output_inferences"]
    assert output["k"]["resolved_output_measure"] == "oscillator_0_100"
    assert output["d"]["resolved_output_measure"] == "oscillator_0_100"
    assert output["j"]["resolved_output_measure"] == "dimensionless"
    assert validation["history_policy"] == "full_history"
    assert r"\mathcal{S}_{3,50}" in output["k"]["display_latex"]
    assert r"\mathcal{W}_{t,9}" in output["k"]["display_latex"]
    assert r"\mathbin{\oslash}_{50}" in output["k"]["display_latex"]
    for item in output.values():
        assert r"\begin{cases}" not in item["display_latex"]
        assert r"\begin{aligned}" not in item["display_latex"]
        assert r"\mathrm{NaN}" not in item["display_latex"]


def test_time_series_excel_export_uses_raw_data_fixed_literals_and_formulas(tmp_path: Path) -> None:
    service, frame = _service(tmp_path)
    artifact = service.export_excel(
        indicator_ids=["builtin-bollinger-bands-series"],
        inline_definition=None,
        targets=[{"kind": "etf", "product_id": "510300.SH"}],
        period="ALL",
    )
    try:
        workbook = load_workbook(artifact.path, data_only=False)
        assert workbook.sheetnames[0] == "01_结果汇总"
        calculation_name = next(name for name in workbook.sheetnames if name.endswith("_计算"))
        result_name = next(name for name in workbook.sheetnames if name.endswith("_结果"))
        calculation = workbook[calculation_name]
        result = workbook[result_name]
        formulas = [
            cell.value
            for row in calculation.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        texts = [str(cell.value) for row in calculation.iter_rows() for cell in row if cell.value is not None]
        assert any("AVERAGE(" in formula for formula in formulas)
        assert any("STDEVP(" in formula for formula in formulas)
        assert any("直接入参 · 收盘价" in text for text in texts)
        assert any("原生 Excel 公式（可复制）" in text for text in texts)
        assert not any("直接入参 · window" in text for text in texts)
        assert not any(
            token in formula
            for formula in formulas
            for token in (
                "AGGREGATE(",
                "SUMPRODUCT(",
                "LOOKUP(",
                "rolling_mean(",
                "rolling_std(",
                "S01_",
                "_xlfn.",
                "_xlws.",
            )
        )
        assert not list(workbook.defined_names)
        assert result["B13"].data_type == "f"
        workbook.close()

        cached = load_workbook(artifact.path, data_only=True)
        cached_result = cached[result_name]
        expected_middle = float(frame["close"].iloc[:20].mean())
        # Result data starts on row 9; a fixed 20-observation window first
        # becomes available on the twentieth row, i.e. Excel row 28.
        assert cached_result["G28"].value == pytest.approx(expected_middle)
        assert cached_result["F28"].value == pytest.approx(cached_result["G28"].value)
        cached.close()
    finally:
        artifact.cleanup()


def test_all_builtin_time_series_indicators_export_formula_workbooks(tmp_path: Path) -> None:
    service, _frame = _service(tmp_path)
    cases = [
        ("builtin-close-moving-average-series", ("AVERAGE(",)),
        ("builtin-bollinger-bands-series", ("AVERAGE(", "STDEVP(")),
        ("builtin-volume-moving-average-series", ("AVERAGE(",)),
        ("builtin-kdj-series", ("MAX(", "MIN(", "ISNUMBER(")),
        (
            "builtin-rolling-5d-annualized-sharpe-series",
            ("AVERAGE(", "STDEV(", "SQRT("),
        ),
    ]
    for indicator_id, formula_tokens in cases:
        artifact = service.export_excel(
            indicator_ids=[indicator_id],
            inline_definition=None,
            targets=[{"kind": "etf", "product_id": "510300.SH"}],
            period="ALL",
        )
        try:
            workbook = load_workbook(artifact.path, data_only=False, read_only=True)
            calculation_name = next(name for name in workbook.sheetnames if name.endswith("_计算"))
            formulas = [
                cell.value
                for row in workbook[calculation_name].iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            joined = "\n".join(formulas)
            assert formulas, indicator_id
            for token in formula_tokens:
                assert token in joined, (indicator_id, token)
            assert not any(
                token in joined
                for token in (
                    "AGGREGATE(",
                    "SUMPRODUCT(",
                    "LOOKUP(",
                    "rolling_mean(",
                    "rolling_std(",
                    "rolling_min(",
                    "rolling_max(",
                    "recursive_smooth(",
                    "divide_or_default(",
                    "S01_",
                    "_xlfn.",
                    "_xlws.",
                )
            )
            assert not list(workbook.defined_names)
            workbook.close()
        finally:
            artifact.cleanup()


def test_time_series_excel_route_returns_xlsx_without_runtime_parameters(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    service, _frame = _service(tmp_path)
    monkeypatch.setattr(custom_indicator_routes, "indicator_service", service)
    app = FastAPI()
    app.include_router(custom_indicator_routes.router)
    client = TestClient(app)
    response = client.post(
        "/api/custom-indicators/export-excel",
        json={
            "indicator_ids": ["builtin-close-moving-average-series"],
            "targets": [{"kind": "etf", "product_id": "510300.SH"}],
            "period": "ALL",
        },
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    with zipfile.ZipFile(io.BytesIO(response.content)) as workbook:
        worksheet_xml = "\n".join(
            workbook.read(name).decode("utf-8")
            for name in workbook.namelist()
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        )
    assert "AVERAGE(" in worksheet_xml
    assert "AGGREGATE(" not in worksheet_xml
    assert "rolling_mean(" not in worksheet_xml
    assert "S01_" not in worksheet_xml
    assert "直接入参 · 收盘价" in worksheet_xml
    assert "直接入参 · window" not in worksheet_xml


def test_scalar_indicator_can_be_lifted_to_fixed_rolling_time_series(
    tmp_path: Path,
) -> None:
    service, _frame = _service(tmp_path)
    generated = service.build_rolling_scalar_draft(
        "builtin-annualized-sharpe-v2",
        1,
        5,
    )
    definition = generated["definition"]
    assert generated["validation"]["valid"] is True
    assert definition["result_kind"] == "time_series"
    assert definition["parameter_schema"] == []
    assert definition["axis_anchor"] == "adjusted_nav"
    assert definition["rolling_source"]["indicator_id"] == "builtin-annualized-sharpe-v2"
    assert definition["rolling_source"]["indicator_revision"] == 1
    assert definition["rolling_source"]["window_observations"] == 5
    expression = definition["series_outputs"][0]["expression"]
    assert "rolling_mean(" in expression
    assert "rolling_std(" in expression
    assert "returns" in expression
    assert "rolling_mean(returns, 5)" in expression
    assert "rolling_std(returns, 5, 1)" in expression
    assert ", 5, 5)" not in expression
    assert ", 5, 1, 5)" not in expression
    assert "mean(returns)" not in expression
    assert "std(returns)" not in expression


def test_rolling_scalar_source_hash_and_formula_are_locked(
    tmp_path: Path,
) -> None:
    service, _frame = _service(tmp_path)
    generated = service.build_rolling_scalar_draft(
        "builtin-annualized-sharpe-v2",
        1,
        5,
    )["definition"]

    wrong_hash = {
        **generated,
        "rolling_source": {
            **generated["rolling_source"],
            "definition_hash": "0" * 64,
        },
    }
    with pytest.raises(ValidationError) as hash_error:
        service.create_indicator(wrong_hash)
    assert hash_error.value.code == "ROLLING_SOURCE_REVISION_MISMATCH"

    changed_formula = {
        **generated,
        "name": "被修改的滚动夏普",
        "series_outputs": [
            {
                **generated["series_outputs"][0],
                "expression": "rolling_mean(returns, 5, 5)",
            }
        ],
    }
    changed_formula["expression"] = changed_formula["series_outputs"][0][
        "expression"
    ]
    with pytest.raises(ValidationError) as formula_error:
        service.create_indicator(changed_formula)
    assert formula_error.value.code == "ROLLING_SOURCE_FORMULA_MISMATCH"


def test_unsupported_scalar_reduction_fails_closed_when_rolling(
    tmp_path: Path,
) -> None:
    service, _frame = _service(tmp_path)
    with pytest.raises(ValidationError) as error:
        service.build_rolling_scalar_draft(
            "builtin-total-return-v2",
            1,
            5,
        )
    assert error.value.code == "ROLLING_SCALAR_OPERATOR_UNSUPPORTED"


def test_five_day_rolling_annualized_sharpe_builtin_is_available_and_njit(
    tmp_path: Path,
) -> None:
    service, _frame = _service(tmp_path)
    definition = service.get_indicator(
        "builtin-rolling-5d-annualized-sharpe-series"
    )
    assert definition["name"] == "5 日滚动年化夏普比率"
    assert definition["rolling_source"]["indicator_id"] == "builtin-annualized-sharpe-v2"
    assert definition["rolling_source"]["window_observations"] == 5
    assert definition["minimum_observations"] == 5
    expression = definition["series_outputs"][0]["expression"]
    assert "rolling_mean(returns, 5)" in expression
    assert "rolling_std(returns, 5, 1)" in expression
    assert ", 5, 5)" not in expression
    assert ", 5, 1, 5)" not in expression
    generated = service.build_rolling_scalar_draft(
        "builtin-annualized-sharpe-v2",
        1,
        5,
    )["definition"]
    generated_expression = generated["series_outputs"][0]["expression"]
    assert generated_expression == expression
    assert generated["rolling_source"]["indicator_id"] == definition[
        "rolling_source"
    ]["indicator_id"]
    assert generated["rolling_source"]["definition_hash"] == definition[
        "rolling_source"
    ]["definition_hash"]

    created = service.create_indicator(generated)
    response = service.evaluate_series(
        indicator_instances=[
            _instance("builtin-rolling-5d-annualized-sharpe-series"),
            _instance(created["id"]),
        ],
        target={"kind": "etf", "product_id": "510300.SH"},
        period="ALL",
    )
    assert response["summary"] == {
        "total": 2,
        "ok": 2,
        "warning": 0,
        "unavailable": 0,
        "error": 0,
    }
    assert response["execution"]["python_fallback"] == 0
    assert response["execution"]["request_time_compilation"] == 0
    results_by_id = {
        item["indicator_id"]: item for item in response["results"]
    }
    builtin_channel = results_by_id[
        "builtin-rolling-5d-annualized-sharpe-series"
    ]["channels"][0]
    custom_channel = results_by_id[created["id"]]["channels"][0]
    assert builtin_channel["id"] == "value"
    assert custom_channel["id"] == "value"
    # Five return observations require six NAV levels; the first aligned
    # return is intentionally null on the NAV date axis.
    assert builtin_channel["values"][:5] == [None, None, None, None, None]
    assert all(value is not None for value in builtin_channel["values"][5:])
    assert _as_float(custom_channel["values"]) == pytest.approx(
        _as_float(builtin_channel["values"]),
        nan_ok=True,
    )


def test_five_day_rolling_sharpe_can_be_selected_for_snapshot_precompute(
    tmp_path: Path,
) -> None:
    service, _frame = _service(tmp_path)
    current = service.get_snapshot_config()
    updated = service.update_snapshot_config(
        current["revision"],
        [
            {
                "indicator_id": "builtin-rolling-5d-annualized-sharpe-series",
                "indicator_revision": 1,
                "period": "1Y",
                "channel_id": "value",
                "reducer": "last_finite",
            }
        ],
    )
    item = updated["items"][0]
    assert item["status"] == "ready"
    assert item["channel_id"] == "value"
    assert item["reducer"] == "last_finite"
    assert item["field"].startswith("metric_")


def test_rolling_scalar_draft_route_returns_locked_provenance(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    service, _frame = _service(tmp_path)
    monkeypatch.setattr(custom_indicator_routes, "indicator_service", service)
    app = FastAPI()
    app.include_router(custom_indicator_routes.router)
    client = TestClient(app)
    response = client.post(
        "/api/custom-indicators/rolling-scalar-draft",
        json={
            "indicator_id": "builtin-annualized-sharpe-v2",
            "indicator_revision": 1,
            "window_observations": 5,
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["source"]["indicator_id"] == "builtin-annualized-sharpe-v2"
    assert payload["source"]["indicator_revision"] == 1
    assert payload["source"]["name"]
    assert payload["definition"]["rolling_source"]["window_observations"] == 5
    assert payload["validation"]["valid"] is True


def test_missing_ohlc_is_unavailable_and_never_filled(tmp_path: Path) -> None:
    frame = _write_market_data(tmp_path)
    frame.drop(columns=["high", "low"]).to_parquet(
        tmp_path / "etf_daily_candle_df.parquet", index=False
    )
    service = CustomIndicatorService(tmp_path, tmp_path)
    service.warm_numba_plans()
    response = service.evaluate_series(
        indicator_instances=[_instance("builtin-kdj-series")],
        target={"kind": "etf", "product_id": "510300.SH"},
        period="ALL",
    )
    assert response["results"][0]["status"] == "unavailable"
    assert response["results"][0]["dates"] == []
    assert response["results"][0]["warnings"][0]["code"] == "VARIABLE_UNAVAILABLE"


def test_output_measure_catalog_and_incompatible_override_are_fail_closed(tmp_path: Path) -> None:
    service, _frame = _service(tmp_path)
    metadata = service.meta()
    catalog = {item["id"] for item in metadata["series_output_measures"]}
    assert metadata["rolling_scalar"] == {
        "supported": True,
        "window_kind": "observations",
        "minimum_window_observations": 2,
        "maximum_window_observations": 5000,
        "transform_version": "1.0.0",
        "draft_endpoint": "/api/custom-indicators/rolling-scalar-draft",
        "source_lock": "indicator_id+revision+definition_hash",
    }
    assert {
        "auto",
        "raw_market_price",
        "adjusted_nav",
        "virtual_nav",
        "normalized",
        "bounded_0_1",
        "bounded_minus1_1",
        "oscillator_0_100",
        "return_decimal",
        "volume",
        "dimensionless",
    } <= catalog

    definition = service.get_indicator("builtin-close-moving-average-series")
    incompatible = {
        **definition,
        "series_outputs": [
            {**definition["series_outputs"][0], "output_measure": "return_decimal"}
        ],
    }
    incompatible["expression"] = incompatible["series_outputs"][0]["expression"]
    validation = service.validate(incompatible)
    assert validation["valid"] is False
    assert validation["diagnostics"][0]["code"] == "SERIES_OUTPUT_MEASURE_MISMATCH"


def test_configuration_argument_cannot_use_data_variable(tmp_path: Path) -> None:
    service, _frame = _service(tmp_path)
    definition = service.get_indicator("builtin-close-moving-average-series")
    invalid = {
        **definition,
        "series_outputs": [
            {
                **definition["series_outputs"][0],
                "expression": "rolling_mean(market_close, observation_count)",
            }
        ],
    }
    invalid["expression"] = invalid["series_outputs"][0]["expression"]
    validation = service.validate(invalid)
    assert validation["valid"] is False
    assert validation["diagnostics"][0]["code"] == "SERIES_CONFIGURATION_MUST_BE_CONSTANT"


def test_scalar_catalog_exposes_rolling_compatibility(tmp_path: Path) -> None:
    service = CustomIndicatorService(tmp_path, tmp_path)
    sharpe = service.get_indicator("builtin-annualized-sharpe-v2")
    drawdown = service.get_indicator("builtin-maximum-drawdown-v2")
    assert sharpe["rolling_series_compatibility"] == {
        "supported": True,
        "protocol_version": "1.0.0",
        "rewritten_reductions": ["mean", "std"],
    }
    assert drawdown["rolling_series_compatibility"]["supported"] is False
    assert drawdown["rolling_series_compatibility"]["code"] == (
        "ROLLING_SCALAR_OPERATOR_UNSUPPORTED"
    )


def test_every_catalogued_rolling_compatible_builtin_derives_successfully(
    tmp_path: Path,
) -> None:
    service = CustomIndicatorService(tmp_path, tmp_path)
    compatible = [
        item
        for item in service.list_indicators()["items"]
        if item.get("source") == "built_in"
        and item.get("rolling_series_compatibility", {}).get("supported")
    ]
    assert {item["id"] for item in compatible} >= {
        "builtin-mean-return-v2",
        "builtin-return-volatility-v2",
        "builtin-annualized-sharpe-v2",
        "builtin-average-volume-v2",
    }
    for source in compatible:
        derived = service.derive_rolling_series(
            indicator_id=source["id"],
            indicator_revision=source["revision"],
            window_observations=20,
        )
        assert derived["validation"]["valid"], source["id"]
        assert derived["definition"]["rolling_source"]["indicator_id"] == source["id"]
        assert derived["definition"]["lookback_observations"] == 20


def test_scalar_indicator_can_be_derived_saved_and_evaluated_as_locked_rolling_series(
    tmp_path: Path,
) -> None:
    service, frame = _service(tmp_path)
    derived = service.derive_rolling_series(
        indicator_id="builtin-annualized-sharpe-v2",
        indicator_revision=1,
        window_observations=10,
    )
    definition = derived["definition"]
    assert derived["validation"]["valid"] is True
    assert definition["name"] == "10 日滚动年化夏普比率"
    assert definition["rolling_source"]["indicator_id"] == "builtin-annualized-sharpe-v2"
    assert definition["rolling_source"]["indicator_revision"] == 1
    assert definition["rolling_source"]["window_observations"] == 10
    assert definition["rolling_source"]["detached"] is False
    assert definition["fixed_parameters"] == [
        {
            "id": "window_observations",
            "label": "滚动观察数",
            "type": "integer",
            "value": 10,
            "source": "rolling_source",
        }
    ]
    assert definition["series_outputs"][0]["expression"] == (
        "(rolling_mean(returns, 10) - risk_free_rate_per_observation) / "
        "rolling_std(returns, 10, 1) * sqrt(periods_per_year)"
    )

    saved = service.create_indicator(definition)
    assert saved["source"] == "custom"
    assert saved["rolling_source"]["definition_hash"] == (
        definition["rolling_source"]["definition_hash"]
    )
    evaluated = service.evaluate_series(
        indicator_instances=[
            {
                "indicator_id": saved["id"],
                "indicator_revision": saved["revision"],
            }
        ],
        target={"kind": "etf", "product_id": "510300.SH"},
        period="ALL",
    )
    values = _as_float(evaluated["results"][0]["channels"][0]["values"])
    returns = pd.Series(frame["adj_nav"]).pct_change()
    risk_free = (1.0 + 0.015) ** (1.0 / 252.0) - 1.0
    expected = (
        (returns.rolling(10, min_periods=10).mean() - risk_free)
        / returns.rolling(10, min_periods=10).std(ddof=1)
        * np.sqrt(252.0)
    ).to_numpy()
    assert values == pytest.approx(expected, nan_ok=True)
    assert evaluated["execution"]["python_fallback"] == 0
    assert evaluated["execution"]["request_time_compilation"] == 0


def test_rolling_source_formula_is_locked_until_explicitly_detached(
    tmp_path: Path,
) -> None:
    service, _frame = _service(tmp_path)
    derived = service.derive_rolling_series(
        indicator_id="builtin-annualized-sharpe-v2",
        indicator_revision=1,
        window_observations=5,
    )["definition"]
    original = derived["series_outputs"][0]["expression"]
    tampered = {
        **derived,
        "expression": f"({original}) + 1",
        "series_outputs": [
            {
                **derived["series_outputs"][0],
                "expression": f"({original}) + 1",
            }
        ],
    }
    invalid = service.validate(tampered)
    assert invalid["valid"] is False
    assert invalid["diagnostics"][0]["code"] == "ROLLING_SOURCE_FORMULA_MISMATCH"

    detached = {
        **tampered,
        "rolling_source": {**tampered["rolling_source"], "detached": True},
    }
    valid = service.validate(detached)
    assert valid["valid"] is True


def test_unsupported_path_dependent_scalar_indicator_fails_closed(
    tmp_path: Path,
) -> None:
    service, _frame = _service(tmp_path)
    with pytest.raises(ValidationError) as error:
        service.derive_rolling_series(
            indicator_id="builtin-maximum-drawdown-v2",
            indicator_revision=1,
            window_observations=20,
        )
    assert error.value.code == "ROLLING_SCALAR_OPERATOR_UNSUPPORTED"


def test_custom_scalar_source_cannot_be_deleted_while_rolling_series_references_it(
    tmp_path: Path,
) -> None:
    service, _frame = _service(tmp_path)
    built_in = service.get_indicator("builtin-annualized-sharpe-v2")
    custom_source = service.create_indicator(
        {**built_in, "name": "工作区年化夏普比率"}
    )
    rolling = service.derive_rolling_series(
        indicator_id=custom_source["id"],
        indicator_revision=custom_source["revision"],
        window_observations=5,
    )["definition"]
    saved_rolling = service.create_indicator(rolling)
    assert saved_rolling["rolling_source"]["indicator_id"] == custom_source["id"]

    with pytest.raises(ConflictError) as error:
        service.delete_indicator(custom_source["id"], custom_source["revision"])
    assert error.value.code == "INDICATOR_IN_ROLLING_SERIES"


def test_derive_rolling_series_route_returns_validated_locked_draft(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    service, _frame = _service(tmp_path)
    monkeypatch.setattr(custom_indicator_routes, "indicator_service", service)
    app = FastAPI()
    app.include_router(custom_indicator_routes.router)
    response = TestClient(app).post(
        "/api/custom-indicators/derive-rolling-series",
        json={
            "indicator_id": "builtin-annualized-sharpe-v2",
            "indicator_revision": 1,
            "window_observations": 5,
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["validation"]["valid"] is True
    assert payload["definition"]["result_kind"] == "time_series"
    assert payload["definition"]["rolling_source"]["indicator_id"] == (
        "builtin-annualized-sharpe-v2"
    )
    assert payload["definition"]["series_outputs"][0]["expression"] == (
        "(rolling_mean(returns, 5) - risk_free_rate_per_observation) / "
        "rolling_std(returns, 5, 1) * sqrt(periods_per_year)"
    )
    latex = payload["validation"]["output_inferences"]["value"]["display_latex"]
    assert r"\mu_{t,5}" in latex
    assert r"s_{t,5}" in latex
    assert r"\sqrt{p_{\mathrm{year}}}" in latex
    assert r"\begin{cases}" not in latex
    assert r"\sum" not in latex
    assert r"\mathrm{NaN}" not in latex
