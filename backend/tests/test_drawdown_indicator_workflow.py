"""Catalog, graph round-trip, selected-output and history integration."""
from __future__ import annotations

import copy

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from custom_indicators.drawdown_indicator import DRAWDOWN_ANALYSIS_ID, drawdown_analysis_builtin
from custom_indicators.errors import NotFoundError
from custom_indicators.service import CustomIndicatorService
from custom_indicators.graph_service import IndicatorGraphService
from custom_indicators.graph_contracts import FormulaResolveRequest, CanvasResolveRequest
from custom_indicators.formula_source import editable_formula_latex, canonical_formula_source
from test_custom_indicator_service import _write_market_data
from test_scalar_outputs_integration import plan_draft


@pytest.fixture
def service(tmp_path):
    _write_market_data(tmp_path)
    owner = CustomIndicatorService(tmp_path, tmp_path)
    owner.scalar_service.warm(owner.get_indicator(DRAWDOWN_ANALYSIS_ID))
    yield owner
    owner.close_compute_engine()


def test_catalog_replaces_scalar_drawdown_and_removes_market_research(service):
    listed = service.list_indicators()["items"]
    ids = {item["id"] for item in listed}
    assert DRAWDOWN_ANALYSIS_ID in ids
    assert "builtin-market-attribution-csi300" not in ids
    assert "builtin-maximum-drawdown-v2" not in ids
    with pytest.raises(NotFoundError):
        service.get_indicator("builtin-market-attribution-csi300")
    assert service.get_indicator("builtin-maximum-drawdown-v2")["output_contract"] == "scalar"
    assert len(service.get_indicator(DRAWDOWN_ANALYSIS_ID)["scalar_outputs"]) == 4


def test_one_call_preview_selection_and_custom_copy(service):
    definition = service.get_indicator(DRAWDOWN_ANALYSIS_ID)
    response = service.evaluate(indicator_ids=[DRAWDOWN_ANALYSIS_ID], inline_definition=None,
                                targets=[{"kind": "etf", "product_id": "510050.SH"}], period="ALL")
    output = response["results"][0]
    assert len(output["outputs"]) == 4
    assert output["outputs"][0]["value"] is not None
    compiled = service.scalar_service._require_warmed(definition).groups[0].compiled
    assert compiled.metadata()["multi_output_call_sites"] == 1
    references = [{"indicator_id": DRAWDOWN_ANALYSIS_ID, "output_id": port} for port in ("max_drawdown", "decline_periods")]
    selected = service.evaluate(indicator_ids=[], inline_definition=None, indicator_refs=references,
                                targets=[{"kind": "etf", "product_id": "510050.SH"}], period="ALL")
    assert len(selected["results"]) == 2
    assert selected["results"][0]["value"] == output["outputs"][0]["value"]
    draft = drawdown_analysis_builtin()
    draft["name"] = "我的回撤分析"
    saved = service.create_indicator(draft)
    assert saved["source"] == "custom"
    changed = copy.deepcopy(saved)
    changed["scalar_outputs"].reverse()
    changed["scalar_outputs"][0]["label"] = "更名后"
    updated = service.update_indicator(saved["id"], saved["revision"], changed)
    assert updated["revision"] == 2
    assert service.indicators.get(saved["id"], 1)["scalar_outputs"][0]["id"] == "max_drawdown"


def declining_data(service):
    path = service.market_data_dir / "etf_daily_df.parquet"
    frame = pd.read_parquet(path)
    mask = frame["ts_code"] == "510050.SH"
    frame.loc[mask, "adj_nav"] = np.linspace(1.0, 0.8, int(mask.sum()))
    frame.to_parquet(path, index=False)


def test_unrecovered_output_does_not_block_max_drawdown_scoring(service):
    declining_data(service)
    definition = service.get_indicator(DRAWDOWN_ANALYSIS_ID)
    response = service.evaluate(indicator_ids=[DRAWDOWN_ANALYSIS_ID], inline_definition=None,
                                targets=[{"kind": "etf", "product_id": "510050.SH"}], period="ALL")
    outputs = {item["output_id"]: item for item in response["results"][0]["outputs"]}
    assert outputs["max_drawdown"]["value"] == pytest.approx(0.2)
    assert outputs["recovery_periods"]["value"] is None
    assert "尚未恢复" in outputs["recovery_periods"]["warnings"][0]["message"]
    plan = service.create_plan(plan_draft(definition, "max_drawdown"))
    result = service.run_plan(plan["id"])
    assert result["ranked_count"] == 1
    assert result["rows"][0]["values"][0]["value"] == pytest.approx(0.2)
    assert result["execution"]["request_time_compilation"] == 0


def test_excel_uses_one_native_scan_sheet_and_leaves_unrecovered_blank(service):
    declining_data(service)
    artifact = service.export_excel(indicator_ids=[DRAWDOWN_ANALYSIS_ID], inline_definition=None,
                                    targets=[{"kind": "etf", "product_id": "510050.SH"}], period="ALL")
    try:
        workbook = load_workbook(artifact.path, data_only=False)
        assert len(workbook.sheetnames) == 2  # one summary, one shared calculation
        summary, calculation = workbook.worksheets
        rows = list(summary.iter_rows(min_row=11, max_row=14, values_only=True))
        assert rows[0][7] == pytest.approx(0.2)
        assert rows[2][7] is None and rows[2][8] is None
        assert "尚未恢复" in rows[2][12]
        assert len({row[8].split("!")[0] for row in rows if isinstance(row[8], str) and row[8].startswith("=")}) == 1
        assert sum(cell.value == "历史峰值" for row in calculation for cell in row) == 1
        formulas = [cell.value for row in calculation for cell in row if cell.data_type == "f"]
        assert any("INDEX(" in formula for formula in formulas)
        assert all("drawdown_analysis" not in formula for formula in formulas)
        assert all("#REF!" not in formula for formula in formulas)
        workbook.close()
    finally:
        artifact.cleanup()


def test_new_snapshot_result_lock_and_legacy_single_result_remain_available(service):
    declining_data(service)
    config = service.get_snapshot_config()
    updated = service.update_snapshot_config(config["revision"], [
        {"indicator_id": DRAWDOWN_ANALYSIS_ID, "indicator_revision": 1, "output_id": "max_drawdown", "period": "ALL"},
        {"indicator_id": "builtin-maximum-drawdown-v2", "indicator_revision": 1, "period": "ALL"},
    ])
    assert all(item["status"] == "ready" for item in updated["items"])
    old = service.evaluate(indicator_ids=["builtin-maximum-drawdown-v2"], inline_definition=None,
                           targets=[{"kind": "etf", "product_id": "510050.SH"}], period="ALL", prefer_snapshot=False)
    assert old["results"][0]["value"] == pytest.approx(0.2)
    assert not old["results"][0].get("outputs")
    current = service.evaluate(indicator_ids=[], inline_definition=None,
        indicator_refs=[{"indicator_id": DRAWDOWN_ANALYSIS_ID, "output_id": "max_drawdown"}],
        targets=[{"kind": "etf", "product_id": "510050.SH"}], period="ALL")
    assert current["results"][0]["value"] == pytest.approx(old["results"][0]["value"])


def test_named_ports_roundtrip_without_fake_projection_nodes(service):
    definition = drawdown_analysis_builtin()
    expressions = {output["id"]: output["expression"] for output in definition["scalar_outputs"]}
    graph_service = IndicatorGraphService(service)
    first = graph_service.resolve(FormulaResolveRequest(source_kind="formula", result_kind="scalar_bundle", expressions=expressions))
    assert first["valid"], first
    nodes = first["graph"]["nodes"]
    assert [node["operator_id"] for node in nodes if node["kind"] == "operator"] == ["drawdown_analysis"]
    assert len(nodes) == 2
    assert {output["port_id"] for output in first["graph"]["outputs"]} == set(expressions)
    second = graph_service.resolve(CanvasResolveRequest(source_kind="graph", result_kind="scalar_bundle", graph=first["graph"]))
    assert second["valid"], second
    assert second["definition_fingerprint"] == first["definition_fingerprint"]
    for expression in expressions.values():
        assert canonical_formula_source(editable_formula_latex(expression)) == expression
    first["graph"]["outputs"][0]["port_id"] = "missing"
    invalid = graph_service.resolve(CanvasResolveRequest(source_kind="graph", result_kind="scalar_bundle", graph=first["graph"]))
    assert not invalid["valid"]


def test_selected_outputs_execute_one_shared_dispatcher_and_reuse_cache(service, monkeypatch):
    definition = service.get_indicator(DRAWDOWN_ANALYSIS_ID)
    compiled = service.scalar_service._require_warmed(definition).groups[0].compiled
    original = type(compiled).compute
    signatures = tuple(compiled.dispatcher.signatures)
    calls = []

    def counted(instance, arguments, enabled):
        if instance is compiled:
            calls.append(instance.plan_id)
        return original(instance, arguments, enabled)

    monkeypatch.setattr(type(compiled), "compute", counted)
    request = dict(indicator_ids=[], inline_definition=None,
        indicator_refs=[{"indicator_id": DRAWDOWN_ANALYSIS_ID, "output_id": output_id}
                        for output_id in ("max_drawdown", "decline_periods", "longest_underwater_periods")],
        targets=[{"kind": "etf", "product_id": "510050.SH"}], period="ALL")
    first = service.evaluate(**request)
    assert len(first["results"]) == 3
    assert calls == [compiled.plan_id]
    second = service.evaluate(**request)
    assert calls == [compiled.plan_id]
    assert second["results"] == first["results"]
    assert tuple(compiled.dispatcher.signatures) == signatures
    assert compiled.metadata()["multi_output_call_sites"] == 1
