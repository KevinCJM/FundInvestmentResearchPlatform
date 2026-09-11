"""Scope safety boundaries and streaming workbook evidence, not only values."""
from __future__ import annotations

import json
import os
from pathlib import Path
import subprocess
import sys

import numpy as np
import pytest
from openpyxl import load_workbook

from cal_indicators.typed_dsl import TypedDslError, compose_typed_series_bundle
from cal_indicators.typed_numba_plan import compile_numba_series_plan, numba_series_plan_id
from cal_indicators import rolling_scope
from test_rolling_interval_graph import build, evaluate
from test_rolling_interval_service import TARGET
import test_rolling_interval_service as rolling_fixtures

service = rolling_fixtures.service


def test_scope_version_is_in_parent_plan_identity(monkeypatch):
    plan = build("mean(market_close)")
    before = numba_series_plan_id(plan)
    monkeypatch.setattr(rolling_scope, "SCOPE_VERSION", "test-distinct-scope")
    assert numba_series_plan_id(plan) != before


def test_dynamic_inner_axes_cannot_reach_unsafe_pairwise_kernel():
    # Some protocols reject this statically; dynamic symbolic axes must also
    # be checked before the two concrete arrays reach correlation's kernel.
    try:
        plan = build("correlation(difference(market_close,1),difference(market_high,2))", "5")
    except TypedDslError as exc:
        assert exc.code in {"SHAPE_MISMATCH", "AXIS_MISMATCH"}
        return
    result = evaluate(plan, {"market_close": np.arange(10.), "market_high": np.arange(10.) ** 2})
    assert np.isnan(result).all()


def test_empty_and_width_one_and_oversize_window():
    x = np.arange(1., 6.)
    np.testing.assert_array_equal(evaluate(build("mean(market_close)", "1"), {"market_close": x}), x)
    assert evaluate(build("mean(market_close)"), {"market_close": x[:0]}).size == 0
    assert np.isnan(evaluate(build("mean(market_close)", "20"), {"market_close": x})).all()


def test_noncontiguous_input_is_rejected_without_compiling_or_copying():
    plan = build("mean(market_close)")
    compiled = compile_numba_series_plan(plan)
    signatures = tuple(compiled.dispatcher.signatures)
    x = np.arange(20.)[::2]
    assert not x.flags.c_contiguous
    with pytest.raises(TypeError):
        evaluate(plan, {"market_close": x})
    assert tuple(compiled.dispatcher.signatures) == signatures


def test_outer_and_inner_uses_of_same_reducer_do_not_change_scope_ownership():
    plan = compose_typed_series_bundle({
        "rolling": "rolling_apply(mean(market_close),3)",
        "scaled": "market_close / mean(market_close)",
    })
    compiled = compile_numba_series_plan(plan)
    x = np.arange(1., 9.)
    context = {"market_close": x, "observation_dates": np.arange(8.) + 20000, "annual_risk_free_rate_decimal": 0.0}
    rolling, scaled = compiled.compute(tuple(context[name] for name in compiled.context_names))
    np.testing.assert_allclose(rolling[2:], [x[i-2:i+1].mean() for i in range(2, len(x))])
    np.testing.assert_allclose(scaled, x / x.mean())


def test_streaming_export_keeps_local_context_formulas_and_valid_cached_values(service):
    custom = service.create_indicator({"name": "区间上下文导出验证", "annual_risk_free_rate_percent": 1.5,
        "expression": "mean(returns)/std(returns,1) + observation_count/observation_count"
        " + window_elapsed_days/window_elapsed_days + risk_free_return_window/risk_free_return_window"})
    saved = service.create_indicator(service.build_rolling_scalar_draft(custom["id"], 1, 3)["definition"])
    artifact = service.export_excel(indicator_ids=[saved["id"]], inline_definition=None, targets=[TARGET], period="1W")
    try:
        workbook = load_workbook(artifact.path, data_only=False)
        local_formulas = {name: [] for name in ("observation_count", "window_elapsed_days", "risk_free_return_window")}
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("直接入参 ·"):
                        for name in local_formulas:
                            if f"（{name}）" in cell.value:
                                local_formulas[name].append(sheet.cell(cell.row + 3, 2).value)
        # The outer context is an input number; each repeated window must have
        # an actual formula, not an overwritten/flushed zero placeholder.
        assert sum(value == "=3" for value in local_formulas["observation_count"]) > 1
        for name in ("window_elapsed_days", "risk_free_return_window"):
            assert sum(isinstance(value, str) and value.startswith("=") for value in local_formulas[name]) > 1
        workbook.close()
        cached = load_workbook(artifact.path, data_only=True)
        for sheet in cached:
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, float):
                        assert np.isfinite(cell.value), (sheet.title, cell.coordinate)
        cached.close()
    finally:
        artifact.cleanup()


def test_njit_window_views_keep_the_original_data_addresses():
    import math
    import numba
    from cal_indicators.typed_operators import get_typed_operator_registry
    plan = build("mean(market_close)")
    scope = next(node for node in plan.nodes if node.operator_id == "rolling_apply")
    capability = rolling_scope.analyze_interval(plan.nodes, scope.inputs[0], get_typed_operator_registry())

    @numba.njit(numba.float64(numba.float64[::1]), nogil=True)
    def address_probe(values):
        # Inspect an address only; never dereference or construct raw pointers.
        return float(values.ctypes.data)

    namespace = {"np": np, "math": math, "interval_body": address_probe}
    exec(rolling_scope._rolling_loop_source(capability.variables, capability), namespace)
    compiled = numba.njit(namespace["rolling_scope"])
    compiled.compile((numba.float64[::1], numba.float64, numba.float64[::1], numba.float64))
    compiled.disable_compile()
    values = np.arange(12., dtype=np.float64)
    actual = compiled(values, 3., np.arange(12.) + 20000., 0.)
    expected = values.ctypes.data + np.arange(10, dtype=np.int64) * values.itemsize
    np.testing.assert_array_equal(actual[2:], expected)


def test_repeated_failures_release_nrt_buffers():
    backend = Path(__file__).resolve().parents[1]
    env = {**os.environ, "NUMBA_NRT_STATS": "1", "OPENBLAS_NUM_THREADS": "1", "OMP_NUM_THREADS": "1", "NUMBA_NUM_THREADS": "1"}
    result = subprocess.run([sys.executable, "scripts/benchmark_rolling_interval.py", "--memory-only"],
                            cwd=backend, env=env, capture_output=True, text=True, timeout=120)
    assert result.returncode == 0, result.stdout + result.stderr
    report = json.loads(result.stdout)
    assert len(report["exception_memory_check"]) >= 6
    assert all(value == 0 for value in report["exception_memory_check"].values())
