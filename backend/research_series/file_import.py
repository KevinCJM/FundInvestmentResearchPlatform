"""Bounded file decoding for user-selected time-series columns; no computation."""

import csv
from datetime import date, datetime
import io
import json
from pathlib import Path
import zipfile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLUMNS = 100


def parse_series_file(content: bytes, filename: str, sheet: str | None = None) -> dict:
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("请选择非空文件，大小不超过 8 MB。")
    suffix = Path(filename).suffix.lower()
    sheets = []
    workbook = None
    try:
        if suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 32 * 1024 * 1024:
                    raise ValueError("Excel 解压后过大，请缩小数据范围。")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            sheets = workbook.sheetnames
            selected = sheet or sheets[0]
            if selected not in sheets:
                raise ValueError("所选工作表不存在。")
            worksheet = workbook[selected]
            # Do not trust a spreadsheet's declared used range.
            worksheet.reset_dimensions()
            iterator = worksheet.iter_rows()
            def excel_rows():
                for cells in iterator:
                    if len(cells) > MAX_COLUMNS:
                        raise ValueError("每个工作表最多支持 100 列。")
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValueError("Excel 含公式，请先复制并粘贴为数值后再上传。")
                    yield [cell.value.isoformat() if isinstance(cell.value, (date, datetime)) else cell.value for cell in cells]
            records = excel_rows()
            columns = next(records, [])
        elif suffix in {".csv", ".json"}:
            try:
                source = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                source = content.decode("gb18030")
            selected = None
            if suffix == ".csv":
                records = csv.reader(io.StringIO(source), strict=True)
                columns = next(records, [])
            else:
                payload = json.loads(source)
                if isinstance(payload, dict):
                    payload = payload.get("rows")
                if not isinstance(payload, list) or not payload or not all(isinstance(row, dict) for row in payload):
                    raise ValueError("JSON 应为非空的行对象数组，或包含 rows 数组。")
                if len(payload) > MAX_ROWS:
                    raise ValueError("单次最多支持 20,000 行。")
                columns = list(dict.fromkeys(key for row in payload for key in row))
                records = ([row.get(key) for key in columns] for row in payload)
        else:
            raise ValueError("支持 CSV、Excel（.xlsx）和 JSON；旧版 .xls 请另存为 .xlsx。")
        columns = [str(value).strip() if value is not None else "" for value in columns]
        if not columns or len(columns) > MAX_COLUMNS or any(not name for name in columns) or len(set(columns)) != len(columns):
            raise ValueError("首行须为不重复的列名，且不超过 100 列。")
        rows = []
        for index, values in enumerate(records):
            if index >= MAX_ROWS:
                raise ValueError("单次最多支持 20,000 行，请拆分文件。")
            if len(values) > len(columns):
                raise ValueError(f"第 {index + 2} 行超出表头列数。")
            if all(value is None or value == "" for value in values):
                continue
            if any(isinstance(value, (dict, list, bool)) for value in values):
                raise ValueError(f"第 {index + 2} 行含不支持的单元格类型。")
            rows.append({column: values[i] if i < len(values) else None for i, column in enumerate(columns)})
        if not rows:
            raise ValueError("文件没有数据行。")
        # Also rejects non-standard JSON NaN/Infinity before returning HTTP JSON.
        json.dumps(rows, allow_nan=False)
        return {"columns": columns, "rows": rows, "sheets": sheets, "sheet": selected}
    except (ValueError, csv.Error, KeyError, UnicodeError, zipfile.BadZipFile, ParseError, InvalidFileException) as exc:
        raise ValueError(str(exc) if isinstance(exc, ValueError) else "文件内容无法读取，请检查格式。") from exc
    finally:
        if workbook is not None:
            workbook.close()
