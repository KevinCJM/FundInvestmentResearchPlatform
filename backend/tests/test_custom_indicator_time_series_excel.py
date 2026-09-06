from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

from custom_indicators.service import CustomIndicatorService


def _write_market_data(data_dir: Path) -> None:
    dates = pd.bdate_range("2026-01-02", periods=80)
    pd.DataFrame(
        [{"ts_code": "510050.SH", "code": "510050", "name": "上证50ETF"}]
    ).to_parquet(data_dir / "etf_info_df.parquet", index=False)
    pd.DataFrame(
        [
            {"ts_code": "510050.SH", "name": "上证50ETF", "date": date, "adj_nav": 1.0 + index * 0.005}
            for index, date in enumerate(dates)
        ]
    ).to_parquet(data_dir / "etf_daily_df.parquet", index=False)
    pd.DataFrame(
        [
            {
                "ts_code": "510050.SH",
                "date": date,
                "open": 2.0 + index * 0.01,
                "high": 2.04 + index * 0.01 + (index % 3) * 0.002,
                "low": 1.96 + index * 0.01 - (index % 2) * 0.002,
                "close": 2.0 + index * 0.01 + ((index % 5) - 2) * 0.003,
                "vol": 100_000.0 + index * 1_000.0,
            }
            for index, date in enumerate(dates)
        ]
    ).to_parquet(data_dir / "etf_daily_candle_df.parquet", index=False)
    pd.DataFrame(
        [{"exchange": "SSE", "cal_date": date.strftime("%Y%m%d"), "is_open": 1} for date in dates]
    ).to_parquet(data_dir / "trade_day_df.parquet", index=False)


def _service(tmp_path: Path) -> CustomIndicatorService:
    _write_market_data(tmp_path)
    service = CustomIndicatorService(tmp_path, tmp_path)
    service.warm_numba_plans()
    return service


def _all_formulas(path: Path) -> list[str]:
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    try:
        return [
            str(cell.value)
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
    finally:
        workbook.close()


def test_time_series_validation_returns_math_latex_measure_and_named_root_dag(tmp_path: Path) -> None:
    service = _service(tmp_path)
    definition = service.get_indicator("builtin-bollinger-bands-series")
    result = service.validate(definition)
    assert result["valid"] is True
    assert set(result["dag"]["roots"]) == {"upper", "middle", "lower"}
    assert set(result["output_inferences"]) == {"upper", "middle", "lower"}
    for channel_id in ("upper", "middle", "lower"):
        output = result["output_inferences"][channel_id]
        assert output["display_latex"]
        assert "rolling_mean" not in output["display_latex"]
        assert output["resolved_output_measure"] == "raw_market_price"
    assert result["lookback_observations"] == 20
    assert all(edge.get("parameter") and isinstance(edge.get("order"), int) for edge in result["dag"]["edges"])
    assert "chart_panel" not in definition
    assert "chart_panel" not in result


def test_time_series_excel_export_uses_raw_data_fixed_literals_and_formulas(tmp_path: Path) -> None:
    service = _service(tmp_path)
    artifact = service.export_excel(
        indicator_ids=["builtin-bollinger-bands-series"],
        inline_definition=None,
        targets=[{"kind": "etf", "product_id": "510050.SH"}],
        period="ALL",
    )
    try:
        formulas = _all_formulas(artifact.path)
        workbook = openpyxl.load_workbook(artifact.path, data_only=False, read_only=False)
        try:
            values = [cell.value for worksheet in workbook.worksheets for row in worksheet.iter_rows() for cell in row]
        finally:
            workbook.close()
        assert any("AVERAGE(" in formula for formula in formulas)
        assert any("STDEVP(" in formula for formula in formulas)
        assert not any(
            token in formula
            for formula in formulas
            for token in (
                "AGGREGATE(",
                "SUMPRODUCT(",
                "LOOKUP(",
                "rolling_mean(",
                "rolling_std(",
                "S01_",
                "_xlfn.",
                "_xlws.",
            )
        )
        assert "原生 Excel 公式（可复制）" in values
        assert "window" not in values
        assert "multiplier" not in values
        text_values = [str(value) for value in values if value is not None]
        assert any("布林上轨" in value for value in text_values)
        assert any("布林中轨" in value for value in text_values)
        assert any("布林下轨" in value for value in text_values)
    finally:
        artifact.cleanup()


def test_five_day_rolling_sharpe_excel_export_uses_native_formulas(
    tmp_path: Path,
) -> None:
    service = _service(tmp_path)
    artifact = service.export_excel(
        indicator_ids=["builtin-rolling-5d-annualized-sharpe-series"],
        inline_definition=None,
        targets=[{"kind": "etf", "product_id": "510050.SH"}],
        period="ALL",
    )
    try:
        formulas = _all_formulas(artifact.path)
        joined = "\n".join(formulas)
        assert "AVERAGE(" in joined
        assert "STDEV(" in joined
        assert "SQRT(" in joined
        assert not any(
            token in joined
            for token in (
                "mean(returns)",
                "std(returns)",
                "rolling_mean(",
                "rolling_std(",
                "S01_",
                "_xlfn.",
                "_xlws.",
            )
        )
    finally:
        artifact.cleanup()


def test_kdj_excel_export_contains_rolling_and_recursive_formulas(tmp_path: Path) -> None:
    service = _service(tmp_path)
    artifact = service.export_excel(
        indicator_ids=["builtin-kdj-series"],
        inline_definition=None,
        targets=[{"kind": "etf", "product_id": "510050.SH"}],
        period="ALL",
    )
    try:
        formulas = _all_formulas(artifact.path)
        assert any("MAX(" in formula for formula in formulas)
        assert any("MIN(" in formula for formula in formulas)
        assert any("IF(ABS(" in formula for formula in formulas)
        assert any("/" in formula and "+" in formula for formula in formulas)
        assert not any(
            token in formula
            for formula in formulas
            for token in (
                "LOOKUP(",
                "AGGREGATE(",
                "SUMPRODUCT(",
                "rolling_min(",
                "rolling_max(",
                "recursive_smooth(",
                "divide_or_default(",
                "S01_",
                "_xlfn.",
                "_xlws.",
            )
        )
    finally:
        artifact.cleanup()
