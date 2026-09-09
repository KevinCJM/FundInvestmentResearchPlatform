"""Named output workbook evidence is version-bound and never fills failures with zero."""
from __future__ import annotations

import openpyxl
import pytest

from test_scalar_indicator_outputs import bundle_draft, service


def test_bundle_excel_has_all_named_results_and_replayable_formulas(service):
    definition = service.create_indicator(bundle_draft())
    kwargs = dict(indicator_ids=[definition["id"]], inline_definition=None,
                  targets=[{"kind": "etf", "product_id": "510050.SH"}], period="ALL")
    result = service.evaluate(**kwargs)["results"][0]
    artifact = service.export_excel(**kwargs)
    try:
        formulas = openpyxl.load_workbook(artifact.path, data_only=False)
        cached = openpyxl.load_workbook(artifact.path, data_only=True)
        assert len(formulas.sheetnames) == 4
        summary = formulas["01_结果汇总"]
        values = cached["01_结果汇总"]
        assert [summary.cell(row, 4).value for row in (11, 12, 13)] == ["mean", "twice", "bad"]
        assert values["H11"].value == pytest.approx(result["outputs"][0]["value"])
        assert values["H12"].value == pytest.approx(result["outputs"][1]["value"])
        assert summary["I11"].data_type == "f"
        assert summary["J11"].data_type == "f"
        assert summary["K11"].data_type == "f"
        assert "%" in summary["H11"].number_format
        assert summary["H13"].value is None
        assert summary["I13"].value is None
        assert summary["K13"].value == "不可比较"
        assert summary["M13"].value
        assert any(cell.data_type == "f" for row in formulas.worksheets[1] for cell in row)
        assert formulas.worksheets[1]["B6"].value == "mean(returns)"
        assert all(summary.cell(row, 5).value == 1 for row in (11, 12, 13))
        formulas.close()
        cached.close()
    finally:
        artifact.cleanup()
    assert not artifact.path.exists()


def test_inline_bundle_export_requires_matching_compile_token(service):
    draft = bundle_draft()
    validation = service.validate(draft)
    artifact = service.export_excel(
        indicator_ids=[], inline_definition=draft, compile_token=validation["compile_token"],
        targets=[{"kind": "fund", "product_id": "000001.OF"}], period="ALL",
    )
    try:
        workbook = openpyxl.load_workbook(artifact.path, data_only=True)
        assert workbook["01_结果汇总"]["H11"].value is not None
        workbook.close()
    finally:
        artifact.cleanup()
